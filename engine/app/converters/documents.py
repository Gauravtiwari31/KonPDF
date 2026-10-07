"""Documents: DOCX, TXT, Markdown and HTML in any direction, plus PDF output.

Every document is read into one small HTML tree (`Node`), then written out
by a renderer: text, Markdown, HTML, DOCX, or PDF through PyMuPDF's HTML
layout engine (`Story`). LibreOffice is used instead when it is installed,
for full Word fidelity; this path is the always-available fallback.
"""

from __future__ import annotations

import html as htmllib
import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

import pymupdf

from ..errors import KonError, Notes

VOID = {"br", "img", "hr", "meta", "link", "input", "col", "source", "wbr"}
SKIP = {"script", "style", "head", "title", "noscript", "template", "svg"}
BLOCKS = {
    "p", "div", "section", "article", "header", "footer", "main", "aside", "nav",
    "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "li", "table", "tr", "pre",
    "blockquote", "hr", "figure", "figcaption", "dl", "dt", "dd", "body", "html",
}  # fmt: skip


@dataclass
class Node:
    tag: str
    attrs: dict[str, str] = field(default_factory=dict)
    children: list["Node | str"] = field(default_factory=list)

    def text(self) -> str:
        return "".join(c if isinstance(c, str) else c.text() for c in self.children)

    def find_all(self, tag: str) -> list["Node"]:
        found = []
        for c in self.children:
            if isinstance(c, Node):
                if c.tag == tag:
                    found.append(c)
                found.extend(c.find_all(tag))
        return found


class _Builder(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]
        self.skipping = 0

    def handle_starttag(self, tag, attrs):
        if tag in SKIP:
            self.skipping += 1
            return
        if self.skipping:
            return
        node = Node(tag, {k: v or "" for k, v in attrs})
        self.stack[-1].children.append(node)
        if tag not in VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        if not self.skipping and tag not in SKIP:
            self.stack[-1].children.append(Node(tag, {k: v or "" for k, v in attrs}))

    def handle_endtag(self, tag):
        if tag in SKIP:
            self.skipping = max(0, self.skipping - 1)
            return
        if self.skipping:
            return
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                break

    def handle_data(self, data):
        if not self.skipping and data:
            self.stack[-1].children.append(data)


def parse_html(source: str) -> Node:
    builder = _Builder()
    builder.feed(source)
    builder.close()
    return builder.root


# ------------------------------------------------------------------ readers


def text_to_html(text: str) -> str:
    paragraphs = re.split(r"\n\s*\n", text.replace("\r\n", "\n").strip())
    return "\n".join(
        "<p>" + "<br>".join(htmllib.escape(line) for line in p.split("\n")) + "</p>" for p in paragraphs if p.strip()
    )


def markdown_to_html(text: str) -> str:
    import markdown

    return markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])


def docx_to_html(path: Path, image_dir: Path, name: str) -> str:
    """Headings, paragraphs (bold/italic/underline), lists, tables and pictures, in order."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(str(path))
    except Exception as e:  # noqa: BLE001
        raise KonError("CORRUPT_FILE", name=name) from e

    image_dir.mkdir(parents=True, exist_ok=True)
    counter = {"n": 0}
    ns_blip = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    ns_embed = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

    def run_html(run) -> str:
        parts = []
        for blip in run.element.iter(ns_blip):
            rid = blip.get(ns_embed)
            part = document.part.related_parts.get(rid) if rid else None
            if part is not None and hasattr(part, "blob"):
                counter["n"] += 1
                ext = (getattr(part, "partname", "") or ".png").rsplit(".", 1)[-1].lower()
                ext = ext if ext in ("png", "jpg", "jpeg", "gif", "bmp") else "png"
                file = f"img{counter['n']}.{ext}"
                (image_dir / file).write_bytes(part.blob)
                parts.append(f'<img src="{file}">')
        text = htmllib.escape(run.text or "")
        if text:
            if run.bold:
                text = f"<b>{text}</b>"
            if run.italic:
                text = f"<i>{text}</i>"
            if run.underline:
                text = f"<u>{text}</u>"
            parts.append(text)
        return "".join(parts)

    def para_html(p: Paragraph) -> tuple[str, str]:
        style = (p.style.name if p.style is not None else "") or ""
        content = "".join(run_html(r) for r in p.runs) or htmllib.escape(p.text)
        is_list = p._p.pPr is not None and p._p.pPr.numPr is not None
        if style.startswith("Heading"):
            level = re.sub(r"\D", "", style) or "2"
            level = str(min(6, max(1, int(level))))
            return "h", f"<h{level}>{content}</h{level}>"
        if style == "Title":
            return "h", f"<h1>{content}</h1>"
        if "List Number" in style:
            return "ol", f"<li>{content}</li>"
        if "List" in style or is_list:
            return "ul", f"<li>{content}</li>"
        align = {1: "center", 2: "right", 3: "justify"}.get(int(p.alignment) if p.alignment is not None else 0)
        style_attr = f' style="text-align:{align}"' if align else ""
        return "p", f"<p{style_attr}>{content}</p>" if content.strip() else "<p>&nbsp;</p>"

    def table_html(t: Table) -> str:
        rows = []
        for r_index, row in enumerate(t.rows):
            cell_tag = "th" if r_index == 0 else "td"
            cells = []
            for cell in row.cells:
                inner = "<br>".join("".join(run_html(r) for r in p.runs) or htmllib.escape(p.text) for p in cell.paragraphs)
                cells.append(f"<{cell_tag}>{inner}</{cell_tag}>")
            rows.append("<tr>" + "".join(cells) + "</tr>")
        return "<table>" + "".join(rows) + "</table>"

    out: list[str] = []
    open_list: str | None = None
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            kind, markup = para_html(Paragraph(child, document))
        elif tag == "tbl":
            kind, markup = "table", table_html(Table(child, document))
        else:
            continue
        if kind in ("ul", "ol"):
            if open_list != kind:
                if open_list:
                    out.append(f"</{open_list}>")
                out.append(f"<{kind}>")
                open_list = kind
        elif open_list:
            out.append(f"</{open_list}>")
            open_list = None
        out.append(markup)
    if open_list:
        out.append(f"</{open_list}>")
    return "\n".join(out)


def read_as_html(path: Path, fmt: str, work: Path, name: str) -> str:
    """Any supported document as an HTML body. Images (from DOCX) go into `work`."""
    if fmt == "docx":
        return docx_to_html(path, work, name)
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    if fmt == "md":
        return markdown_to_html(text)
    if fmt == "html":
        return text
    return text_to_html(text)


# ---------------------------------------------------------------- renderers


def _inline_text(node: Node | str) -> str:
    if isinstance(node, str):
        return re.sub(r"\s+", " ", node)
    if node.tag == "br":
        return "\n"
    if node.tag == "img":
        return ""
    return "".join(_inline_text(c) for c in node.children)


def to_text(root: Node) -> str:
    lines: list[str] = []

    def walk(node: Node | str, depth: int = 0) -> None:
        if isinstance(node, str):
            if node.strip():
                lines.append(re.sub(r"\s+", " ", node).strip())
            return
        if node.tag in ("ul", "ol"):
            for i, li in enumerate(c for c in node.children if isinstance(c, Node) and c.tag == "li"):
                bullet = f"{i + 1}." if node.tag == "ol" else "•"
                lines.append("  " * depth + f"{bullet} {_inline_text(li).strip()}")
            lines.append("")
            return
        if node.tag == "table":
            for tr in node.find_all("tr"):
                cells = [_inline_text(c).strip() for c in tr.children if isinstance(c, Node) and c.tag in ("td", "th")]
                lines.append("\t".join(cells))
            lines.append("")
            return
        if node.tag in BLOCKS - {"body", "html", "div", "section", "article", "main"} or node.tag.startswith("h"):
            text = _inline_text(node).strip()
            if text:
                lines.append(text)
                lines.append("")
            return
        if node.tag == "hr":
            lines.append("")
            return
        for c in node.children:
            walk(c, depth)

    walk(root)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"


def _inline_md(node: Node | str) -> str:
    if isinstance(node, str):
        return re.sub(r"\s+", " ", node)
    inner = "".join(_inline_md(c) for c in node.children)
    if node.tag in ("b", "strong") and inner.strip():
        return f"**{inner.strip()}** "
    if node.tag in ("i", "em") and inner.strip():
        return f"*{inner.strip()}* "
    if node.tag == "code":
        return f"`{inner}`"
    if node.tag == "a" and node.attrs.get("href"):
        return f"[{inner.strip()}]({node.attrs['href']})"
    if node.tag == "br":
        return "  \n"
    return inner


def to_markdown(root: Node) -> str:
    out: list[str] = []

    def walk(node: Node | str) -> None:
        if isinstance(node, str):
            if node.strip():
                out.append(node.strip())
            return
        tag = node.tag
        if re.fullmatch(r"h[1-6]", tag):
            out.append("#" * int(tag[1]) + " " + _inline_md(node).strip())
        elif tag == "p":
            text = _inline_md(node).strip()
            if text:
                out.append(text)
        elif tag in ("ul", "ol"):
            items = [c for c in node.children if isinstance(c, Node) and c.tag == "li"]
            marker = (lambda i: f"{i + 1}. ") if tag == "ol" else (lambda i: "- ")
            out.append("\n".join(marker(i) + _inline_md(li).strip() for i, li in enumerate(items)))
        elif tag == "table":
            rows = [
                [_inline_md(c).strip().replace("|", "\\|") for c in tr.children if isinstance(c, Node) and c.tag in ("td", "th")]
                for tr in node.find_all("tr")
            ]
            rows = [r for r in rows if r]
            if rows:
                width = max(len(r) for r in rows)
                rows = [r + [""] * (width - len(r)) for r in rows]
                out.append("\n".join(["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width] + ["| " + " | ".join(r) + " |" for r in rows[1:]]))
        elif tag == "pre":
            out.append("```\n" + node.text().strip("\n") + "\n```")
        elif tag == "blockquote":
            out.append("> " + _inline_md(node).strip())
        elif tag == "hr":
            out.append("---")
        else:
            for c in node.children:
                walk(c)

    walk(root)
    return "\n\n".join(x for x in out if x) + "\n"


def to_html_document(body: str, title: str) -> str:
    return (
        "<!doctype html>\n<html><head><meta charset=\"utf-8\">"
        f"<title>{htmllib.escape(title)}</title><style>{READ_CSS}</style></head>"
        f"<body>\n{body}\n</body></html>\n"
    )


READ_CSS = (
    "body{font-family:system-ui,sans-serif;max-width:46rem;margin:2rem auto;padding:0 1rem;line-height:1.5}"
    "table{border-collapse:collapse}td,th{border:1px solid #999;padding:4px 8px}img{max-width:100%}"
)

PDF_CSS = """
* { font-family: sans-serif; }
body { font-size: 11pt; line-height: 1.4; }
h1 { font-size: 22pt; margin: 0 0 10pt 0; }
h2 { font-size: 17pt; margin: 14pt 0 8pt 0; }
h3 { font-size: 14pt; margin: 12pt 0 6pt 0; }
h4, h5, h6 { font-size: 12pt; margin: 10pt 0 6pt 0; }
p { margin: 0 0 8pt 0; }
table { border-collapse: collapse; margin: 6pt 0 10pt 0; }
th, td { border: 0.6pt solid #888; padding: 3pt 5pt; font-size: 9.5pt; }
th { background-color: #EEF0F8; font-weight: bold; }
pre, code { font-family: monospace; font-size: 9.5pt; }
li { margin-bottom: 3pt; }
"""


def html_to_pdf(body: str, out: Path, archive_dir: Path | None = None, landscape: bool = False) -> Path:
    """Lays HTML out on A4 pages with PyMuPDF's Story engine."""
    archive = pymupdf.Archive(str(archive_dir)) if archive_dir and archive_dir.is_dir() else None
    story = pymupdf.Story(html=body, user_css=PDF_CSS, archive=archive)
    mediabox = pymupdf.paper_rect("a4-l" if landscape else "a4")
    where = mediabox + (50, 50, -50, -50)
    writer = pymupdf.DocumentWriter(str(out))
    more = 1
    pages = 0
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
        pages += 1
        if pages > 2000:  # a runaway layout must not run forever
            break
    writer.close()
    return out


def to_docx(root: Node, out: Path, image_dir: Path | None = None) -> Path:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    section = document.sections[0]
    usable = (section.page_width - section.left_margin - section.right_margin) / 914400

    def add_inline(paragraph, node: Node | str, bold=False, italic=False, underline=False) -> None:
        if isinstance(node, str):
            text = re.sub(r"\s+", " ", node)
            if text:
                run = paragraph.add_run(text)
                run.bold, run.italic, run.underline = bold or None, italic or None, underline or None
            return
        if node.tag == "br":
            paragraph.add_run().add_break()
            return
        if node.tag == "img":
            src = node.attrs.get("src", "")
            path = (image_dir / Path(src).name) if image_dir and src and not src.startswith(("http:", "https:", "data:")) else None
            if path and path.is_file():
                try:
                    paragraph.add_run().add_picture(str(path), width=Inches(min(usable, 6)))
                except Exception:  # noqa: BLE001 - unsupported image type is skipped
                    pass
            return
        b = bold or node.tag in ("b", "strong", "th")
        i = italic or node.tag in ("i", "em")
        u = underline or node.tag == "u"
        for c in node.children:
            add_inline(paragraph, c, b, i, u)

    def walk(node: Node | str) -> None:
        if isinstance(node, str):
            if node.strip():
                document.add_paragraph(node.strip())
            return
        tag = node.tag
        if re.fullmatch(r"h[1-6]", tag):
            heading = document.add_heading(level=int(tag[1]))
            add_inline(heading, node)
        elif tag in ("p", "blockquote", "figcaption", "dt", "dd"):
            add_inline(document.add_paragraph(), node)
        elif tag == "pre":
            p = document.add_paragraph()
            p.add_run(node.text()).font.name = "Courier New"
        elif tag in ("ul", "ol"):
            style = "List Number" if tag == "ol" else "List Bullet"
            for li in (c for c in node.children if isinstance(c, Node) and c.tag == "li"):
                add_inline(document.add_paragraph(style=style), li)
        elif tag == "table":
            rows = [[c for c in tr.children if isinstance(c, Node) and c.tag in ("td", "th")] for tr in node.find_all("tr")]
            rows = [r for r in rows if r]
            if rows:
                width = max(len(r) for r in rows)
                table = document.add_table(rows=len(rows), cols=width)
                table.style = "Table Grid"
                for r_index, row in enumerate(rows):
                    for c_index, cell in enumerate(row):
                        target = table.cell(r_index, c_index).paragraphs[0]
                        add_inline(target, cell, bold=cell.tag == "th")
        elif tag == "hr":
            document.add_page_break()
        elif tag == "img":
            add_inline(document.add_paragraph(), node)
        else:
            inline_only = all(isinstance(c, str) or c.tag not in BLOCKS for c in node.children)
            if inline_only and node.tag not in ("root", "body", "html") and _inline_text(node).strip():
                add_inline(document.add_paragraph(), node)
            else:
                for c in node.children:
                    walk(c)

    walk(root)
    document.save(out)
    return out


def html_tables(body: str, name: str):
    """Tables in an HTML body as {sheet name: DataFrame}."""
    import pandas as pd

    try:
        frames = pd.read_html(io.StringIO(body))
    except ValueError:
        frames = []
    if not frames:
        raise KonError("NO_TABLES", name=name)
    return {f"Table {i + 1}": frame for i, frame in enumerate(frames)}
