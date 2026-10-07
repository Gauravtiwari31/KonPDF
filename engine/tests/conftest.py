"""Shared fixtures: real sample files of every kind, made on the fly."""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
from pathlib import Path

import pytest

# Isolated data folder and no rate limit for tests; set before the app is imported.
os.environ.setdefault("KON_DATA_DIR", tempfile.mkdtemp(prefix="konpdf-test-"))
os.environ.setdefault("KON_RATE_LIMIT", "0")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pymupdf  # noqa: E402
from PIL import Image, ImageDraw  # noqa: E402


def _photo(size=(1600, 1200)) -> Image.Image:
    """A noisy gradient: compresses like a real photo, not like a flat colour."""
    import numpy as np

    w, h = size
    x = np.linspace(0, 255, w, dtype=np.float32)
    y = np.linspace(0, 255, h, dtype=np.float32)[:, None]
    rng = np.random.default_rng(7)
    r = (x + y) / 2 + rng.normal(0, 18, (h, w))
    g = np.abs(x - y) + rng.normal(0, 18, (h, w))
    b = 255 - (x + y) / 2 + rng.normal(0, 18, (h, w))
    arr = np.clip(np.dstack([r, g, b]), 0, 255).astype("uint8")
    img = Image.fromarray(arr, "RGB")
    ImageDraw.Draw(img).rectangle([w // 4, h // 4, w // 2, h // 2], outline="white", width=12)
    return img


@pytest.fixture(scope="session")
def samples(tmp_path_factory) -> dict[str, Path]:
    d = tmp_path_factory.mktemp("samples")
    out: dict[str, Path] = {}

    photo = _photo()
    out["jpg"] = d / "photo.jpg"
    photo.save(out["jpg"], "JPEG", quality=95, dpi=(72, 72))

    logo = Image.new("RGBA", (400, 300), (0, 0, 0, 0))
    ImageDraw.Draw(logo).ellipse([50, 30, 350, 270], fill=(115, 87, 255, 255))
    out["png"] = d / "logo.png"
    logo.save(out["png"])

    out["webp"] = d / "photo.webp"
    photo.resize((800, 600)).save(out["webp"], "WEBP", quality=80)

    out["gif"] = d / "anim.gif"
    frames = [Image.new("RGB", (120, 80), c) for c in ("red", "green", "blue")]
    frames[0].save(out["gif"], save_all=True, append_images=frames[1:], duration=100, loop=0)

    out["tiff"] = d / "scan.tiff"
    pages = [Image.new("RGB", (300, 400), "white") for _ in range(2)]
    pages[0].save(out["tiff"], save_all=True, append_images=pages[1:])

    out["svg"] = d / "shape.svg"
    out["svg"].write_text('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100"><rect width="200" height="100" fill="#7357FF"/></svg>')

    try:
        from pillow_heif import register_heif_opener

        register_heif_opener()
        out["heic"] = d / "iphone.heic"
        photo.resize((640, 480)).save(out["heic"], "HEIF", quality=80)
    except Exception:  # noqa: BLE001 - HEIC encoder may be missing; tests skip it
        pass

    # A 3-page PDF with real text, a heading and a table.
    pdf = pymupdf.open()
    for n in range(3):
        page = pdf.new_page()
        page.insert_text((72, 80), f"Chapter {n + 1}", fontsize=22)
        page.insert_text((72, 120), f"This is page {n + 1} of a KonPDF test document with real text.", fontsize=11)
    page = pdf[0]
    x0, y0 = 72, 200
    for r, row in enumerate([["Name", "Qty", "Price"], ["Apple", "3", "1.20"], ["Pear", "5", "0.80"]]):
        for c, value in enumerate(row):
            rect = pymupdf.Rect(x0 + c * 120, y0 + r * 24, x0 + (c + 1) * 120, y0 + (r + 1) * 24)
            page.draw_rect(rect, color=(0, 0, 0), width=0.8)
            page.insert_text((rect.x0 + 6, rect.y1 - 7), value, fontsize=10)
    out["pdf"] = d / "report.pdf"
    pdf.save(out["pdf"])

    # A "scanned" PDF: one big image, no text.
    scanned = pymupdf.open()
    buf = io.BytesIO()
    _photo((800, 1100)).save(buf, "JPEG", quality=85)
    scanned.new_page().insert_image(pymupdf.Rect(0, 0, 595, 842), stream=buf.getvalue())
    out["scanned_pdf"] = d / "scan.pdf"
    scanned.save(out["scanned_pdf"])

    from docx import Document

    doc = Document()
    doc.add_heading("KonPDF test", level=1)
    para = doc.add_paragraph("Hello ")
    para.add_run("bold").bold = True
    para.add_run(" and ")
    para.add_run("italic").italic = True
    doc.add_paragraph("First point", style="List Bullet")
    doc.add_paragraph("Second point", style="List Bullet")
    table = doc.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("City", "Pop"), ("Pune", "7"), ("Delhi", "32")]):
        table.cell(r, 0).text, table.cell(r, 1).text = a, b
    pic = io.BytesIO()
    photo.resize((320, 240)).save(pic, "PNG")
    doc.add_picture(pic)
    doc.add_paragraph("नमस्ते दुनिया")  # Hindi text must survive
    out["docx"] = d / "letter.docx"
    doc.save(out["docx"])

    out["txt"] = d / "notes.txt"
    out["txt"].write_text("First line\nsecond line\n\nNew paragraph.", "utf-8")
    out["md"] = d / "readme.md"
    out["md"].write_text("# Title\n\nSome **bold** text.\n\n- one\n- two\n\n| a | b |\n|---|---|\n| 1 | 2 |\n", "utf-8")
    out["html"] = d / "page.html"
    out["html"].write_text("<html><body><h1>Hi</h1><p>Para</p><table><tr><th>x</th><th>y</th></tr><tr><td>1</td><td>2</td></tr></table></body></html>", "utf-8")

    import pandas as pd

    df = pd.DataFrame({"Name": ["Asha", "Ravi", "Zoë"], "Score": [91, 78, 88]})
    out["xlsx"] = d / "marks.xlsx"
    with pd.ExcelWriter(out["xlsx"]) as w:
        df.to_excel(w, sheet_name="Class A", index=False)
        df.to_excel(w, sheet_name="Class B", index=False)
    out["csv"] = d / "marks.csv"
    df.to_csv(out["csv"], index=False, sep=";")
    out["json"] = d / "marks.json"
    out["json"].write_text(json.dumps(df.to_dict(orient="records"), ensure_ascii=False), "utf-8")

    out["mp4"] = d / "clip.mp4"
    out["mp4"].write_bytes(b"\x00\x00\x00\x18ftypmp42" + b"\x00" * 64)
    out["empty"] = d / "empty.pdf"
    out["empty"].write_bytes(b"")
    out["fake_pdf"] = d / "fake.pdf"
    out["fake_pdf"].write_bytes(b"%PDF-1.4\nthis is not really a pdf")
    return out


@pytest.fixture()
def client():
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as c:
        yield c


@pytest.fixture()
def workdirs(tmp_path):
    out, work = tmp_path / "out", tmp_path / "work"
    out.mkdir()
    work.mkdir()
    return out, work
